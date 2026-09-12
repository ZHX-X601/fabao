"""
文书生成接口
"""

from io import BytesIO

from docx import Document as DocxDocument
from docx.shared import Pt
from fastapi import APIRouter, Depends
from fastapi.responses import StreamingResponse
from pydantic import BaseModel
from sqlalchemy.orm import Session

from app.core.dependencies import get_current_user, get_db
from app.core.exceptions import NotFoundError
from app.models.document import Document
from app.models.user import User
from app.schemas.document import DocumentGenerateRequest, DocumentOut
from app.services.document_service import generate_document
from app.utils.response import ok

router = APIRouter(prefix="/documents")


class DownloadRequest(BaseModel):
    """文书下载请求体：直接传入已生成的文书内容与类型"""

    content: str  # 文书全文
    doc_type: str  # 文书类型（用于文件名）


@router.post("/generate", summary="生成法律文书")
async def generate(
    body: DocumentGenerateRequest,
    current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    """
    提交表单数据生成文书

    流程：模板拼接生成全文 -> 保存生成记录 -> 返回记录
    """
    content = generate_document(body.doc_type, body.form_data)
    document = Document(
        user_id=current_user.id,
        doc_type=body.doc_type,
        form_data=body.form_data,
        generated_content=content,
    )
    db.add(document)
    db.commit()
    db.refresh(document)
    return ok(DocumentOut.model_validate(document))


@router.get("/list", summary="获取当前用户的文书历史列表")
def list_documents(
    current_user: User = Depends(get_current_user), db: Session = Depends(get_db)
):
    """按创建时间倒序返回当前用户的全部文书记录"""
    documents = (
        db.query(Document)
        .filter(Document.user_id == current_user.id)
        .order_by(Document.created_at.desc())
        .all()
    )
    return ok([DocumentOut.model_validate(d) for d in documents])


@router.get("/{document_id}", summary="获取某条文书记录")
def get_document(
    document_id: int,
    current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    """
    获取指定文书记录详情

    注意：本路由必须声明在 /list 之后，否则 "list" 会被误解析为 id 参数
    :raises NotFoundError: 记录不存在或不属于当前用户
    """
    document = (
        db.query(Document)
        .filter(Document.id == document_id, Document.user_id == current_user.id)
        .first()
    )
    if document is None:
        raise NotFoundError("文书记录不存在")
    return ok(DocumentOut.model_validate(document))


@router.post("/download", summary="下载文书为 Word 文件")
async def download_document(
    body: DownloadRequest,
    current_user: User = Depends(get_current_user),
):
    """
    将文书内容生成为 Word（.docx）文件并返回下载

    流程：用 python-docx 把文书全文写入内存中的 .docx -> 以流方式返回
    """
    # 创建 Word 文档
    doc = DocxDocument()

    # 设置正文默认字体
    style = doc.styles["Normal"]
    style.font.name = "宋体"
    style.font.size = Pt(12)

    # 按行写入文书内容（空行用空段落表示，保留原文格式）
    for line in body.content.split("\n"):
        paragraph = doc.add_paragraph(line)
        # 标题行（不含冒号且较短的行）加粗，模拟文书标题效果
        if line and len(line) <= 12 and "：" not in line and "，" not in line:
            for run in paragraph.runs:
                run.bold = True

    # 写入内存缓冲区
    buffer = BytesIO()
    doc.save(buffer)
    buffer.seek(0)

    # 文件名：文书类型 + 时间戳（中文需 URL 编码，HTTP header 只支持 latin-1）
    from datetime import datetime
    from urllib.parse import quote

    filename = f"{body.doc_type}_{datetime.now().strftime('%Y%m%d%H%M%S')}.docx"
    encoded_filename = quote(filename)

    # 以流方式返回 Word 文件
    return StreamingResponse(
        buffer,
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        headers={"Content-Disposition": f"attachment; filename*=UTF-8''{encoded_filename}"},
    )
